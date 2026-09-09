# app/sources.py
"""Reading the two inputs: a sitemap (URL or file) and a list of dead URLs."""

from __future__ import annotations

import csv
import gzip
import io
import ipaddress
import os
import re
import socket
from urllib.parse import urlparse

import httpx
import openpyxl

LOC_RE = re.compile(r"<loc>\s*(.*?)\s*</loc>", re.I | re.S)
URL_RE = re.compile(r"https?://[^\s\"'<>,;]+", re.I)

USER_AGENT = "Mozilla/5.0 (compatible; RedirectMapper/1.0; +sitemap-to-404-matcher)"
MAX_SITEMAP_FILES = 60


def is_public_url(url: str) -> bool:
    """
    Reject anything that would make the server fetch its own network.
    Set ALLOW_PRIVATE_SITEMAPS=1 when running against a staging host on a LAN.
    """
    if os.environ.get("ALLOW_PRIVATE_SITEMAPS") == "1":
        return True
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        return False
    host = parsed.hostname.lower()
    if host in ("localhost", "metadata.google.internal") or host.endswith(".local"):
        return False
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        return False
    for info in infos:
        try:
            ip = ipaddress.ip_address(info[4][0])
        except ValueError:
            return False
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
            return False
    return True


def _decode(body: bytes) -> str:
    if body[:2] == b"\x1f\x8b":
        body = gzip.decompress(body)
    return body.decode("utf-8", "replace")


def _unescape(url: str) -> str:
    return (
        url.replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&#39;", "'")
        .strip()
    )


def parse_sitemap_text(text: str) -> tuple[list[str], list[str]]:
    """Return (page urls, child sitemap urls) from one sitemap document."""
    locs = [_unescape(u) for u in LOC_RE.findall(text)]
    if not locs:
        # Plain-text sitemaps are one URL per line.
        locs = [line.strip() for line in text.splitlines() if line.strip().startswith("http")]

    is_index = "<sitemapindex" in text.lower()
    if is_index:
        return [], locs
    # Mixed or malformed files: treat nested .xml entries as children.
    pages, children = [], []
    for u in locs:
        (children if u.lower().split("?")[0].endswith((".xml", ".xml.gz")) else pages).append(u)
    return pages, children


async def load_sitemap_from_urls(urls: list[str], progress=None) -> tuple[list[str], list[str]]:
    """Fetch one or more sitemaps, following index files. Returns (urls, log lines)."""
    log: list[str] = []
    found: list[str] = []
    queue = [u.strip() for u in urls if u.strip()]
    seen: set[str] = set()

    limits = httpx.Limits(max_connections=8)
    async with httpx.AsyncClient(
        follow_redirects=True, timeout=45.0, headers={"User-Agent": USER_AGENT}, limits=limits
    ) as client:
        while queue and len(seen) < MAX_SITEMAP_FILES:
            target = queue.pop(0)
            if target in seen:
                continue
            seen.add(target)
            if not is_public_url(target):
                log.append(f"refused {target} — not a public address")
                continue
            try:
                resp = await client.get(target)
                resp.raise_for_status()
            except Exception as exc:  # network, DNS, 4xx, 5xx
                log.append(f"could not read {target} — {type(exc).__name__}")
                continue

            pages, children = parse_sitemap_text(_decode(resp.content))
            found.extend(pages)
            queue.extend(c for c in children if c not in seen)
            if children and not pages:
                log.append(f"{target}: index with {len(children)} child sitemaps")
            else:
                log.append(f"{target}: {len(pages)} URLs")
            if progress:
                progress(len(found))

    return _dedupe(found), log


def load_sitemap_from_files(files: list[tuple[str, bytes]]) -> tuple[list[str], list[str]]:
    log: list[str] = []
    found: list[str] = []
    for name, blob in files:
        try:
            pages, children = parse_sitemap_text(_decode(blob))
        except Exception as exc:
            log.append(f"could not read {name} — {type(exc).__name__}")
            continue
        found.extend(pages)
        if children and not pages:
            log.append(
                f"{name}: this is a sitemap index listing {len(children)} child files — "
                "paste the sitemap URL instead so they can be followed"
            )
        else:
            log.append(f"{name}: {len(pages)} URLs")
    return _dedupe(found), log


def _dedupe(urls: list[str]) -> list[str]:
    out, seen = [], set()
    for u in urls:
        u = u.strip()
        if not u or not u.lower().startswith("http"):
            continue
        if u.lower().split("?")[0].endswith((".xml", ".xml.gz")):
            continue
        if u not in seen:
            seen.add(u)
            out.append(u)
    return out


# --------------------------------------------------------------------------
# dead URL list
# --------------------------------------------------------------------------

def _looks_like_url(value) -> bool:
    return isinstance(value, str) and value.strip().lower().startswith("http")


def load_dead_urls(name: str, blob: bytes) -> tuple[list[str], str]:
    """Pull URLs out of an xlsx, csv or txt export. Returns (urls, description)."""
    lower = name.lower()

    if lower.endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return [], "the sheet is empty"

        width = max(len(r) for r in rows)
        best_col, best_hits = 0, 0
        for col in range(width):
            hits = sum(1 for r in rows if col < len(r) and _looks_like_url(r[col]))
            if hits > best_hits:
                best_col, best_hits = col, hits
        if not best_hits:
            return [], "no column in the first sheet contains URLs"

        header = rows[0][best_col] if best_col < len(rows[0]) else None
        urls = [
            r[best_col].strip()
            for r in rows
            if best_col < len(r) and _looks_like_url(r[best_col])
        ]
        label = f'column "{header}"' if isinstance(header, str) else f"column {best_col + 1}"
        return _dedupe_keep_order(urls), f"{len(urls)} URLs from {label}"

    text = blob.decode("utf-8", "replace")

    if lower.endswith((".csv", ".tsv")):
        delim = "\t" if lower.endswith(".tsv") else ","
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))
        urls = [cell.strip() for row in rows for cell in row if _looks_like_url(cell)]
        return _dedupe_keep_order(urls), f"{len(urls)} URLs from {name}"

    urls = [m.group(0) for m in URL_RE.finditer(text)]
    return _dedupe_keep_order(urls), f"{len(urls)} URLs from {name}"


def _dedupe_keep_order(urls: list[str]) -> list[str]:
    out, seen = [], set()
    for u in urls:
        if u not in seen:
            seen.add(u)
            out.append(u)
    return out


def same_host(a: str, b: str) -> bool:
    ha = urlparse(a).netloc.lower().removeprefix("www.")
    hb = urlparse(b).netloc.lower().removeprefix("www.")
    return ha == hb
