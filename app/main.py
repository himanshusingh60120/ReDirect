# app/main.py
"""
Redirect Mapper — connect a list of 404s to their closest live URL.

Stateless by design: a match returns the rows to the browser, and the browser
posts them back when you download. Nothing is held between requests, so this
behaves the same on one long-lived process or on serverless functions that
never see each other.

Run locally:
    uvicorn app.main:app --reload
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from .matcher import Index
from .sources import (
    load_dead_urls,
    load_sitemap_from_files,
    load_sitemap_from_urls,
    same_host,
)

APP_DIR = Path(__file__).parent
STATIC_DIR = APP_DIR / "static"

# Vercel rejects request bodies over 4.5 MB before they reach the function,
# so uploads are capped below that and large sitemaps go in by URL instead.
MAX_UPLOAD_BYTES = 4 * 1024 * 1024

app = FastAPI(title="Redirect Mapper", docs_url="/api/docs", redoc_url=None)


class Row(BaseModel):
    dead: str = ""
    live: str = ""
    score: float = 0.0
    method: str = ""
    matched_on: str = ""
    notes: str = ""


class ExportRequest(BaseModel):
    rows: list[Row] = Field(default_factory=list)


@app.get("/health")
async def health():
    return {"ok": True}


# --------------------------------------------------------------------------
# matching
# --------------------------------------------------------------------------

@app.post("/api/match")
async def match(
    sitemap_urls: str = Form(""),
    sitemap_files: list[UploadFile] = File(default=[]),
    dead_file: UploadFile | None = File(default=None),
    dead_text: str = Form(""),
):
    log: list[str] = []
    live: list[str] = []

    urls = [u.strip() for u in sitemap_urls.replace(",", "\n").splitlines() if u.strip()]
    if urls:
        fetched, notes = await load_sitemap_from_urls(urls)
        live.extend(fetched)
        log.extend(notes)

    uploaded = []
    for f in sitemap_files or []:
        blob = await f.read()
        if not blob:
            continue
        if len(blob) > MAX_UPLOAD_BYTES:
            raise HTTPException(
                413,
                f"{f.filename} is over 4 MB. Paste the sitemap URL instead — "
                "it is read server-side with no size limit.",
            )
        uploaded.append((f.filename or "sitemap.xml", blob))
    if uploaded:
        parsed, notes = load_sitemap_from_files(uploaded)
        live.extend(parsed)
        log.extend(notes)

    live = list(dict.fromkeys(live))
    if not live:
        raise HTTPException(
            400,
            "No live URLs found. Paste a sitemap URL (an index file works — its "
            "children are followed) or upload the XML.",
        )

    dead: list[str] = []
    if dead_file is not None and dead_file.filename:
        blob = await dead_file.read()
        if len(blob) > MAX_UPLOAD_BYTES:
            raise HTTPException(413, f"{dead_file.filename} is over 4 MB.")
        dead, note = load_dead_urls(dead_file.filename, blob)
        log.append(note)
    if dead_text.strip():
        extra, note = load_dead_urls("pasted.txt", dead_text.encode())
        seen = set(dead)
        dead.extend(u for u in extra if u not in seen)
        log.append(note)

    if not dead:
        raise HTTPException(400, "No URLs found in the 404 list. Upload a CSV, XLSX or plain list.")

    index = Index(live)
    rows = []
    for url in dead:
        result = index.match(url)
        rows.append(
            {
                "dead": url,
                "live": result["url"],
                "score": result["score"],
                "method": result["method"],
                "matched_on": result["matched_on"],
                "notes": "; ".join(result["notes"]),
                "alternates": result["alternates"],
            }
        )

    off_host = sum(1 for r in rows if r["live"] and not same_host(r["dead"], r["live"]))
    if off_host:
        log.append(f"{off_host} matches point at a different hostname than the 404 — worth a look")

    strong = sum(1 for r in rows if r["score"] >= 85)
    fair = sum(1 for r in rows if 60 <= r["score"] < 85)

    return JSONResponse(
        {
            "live_count": len(live),
            "dead_count": len(rows),
            "strong": strong,
            "fair": fair,
            "weak": len(rows) - strong - fair,
            "log": log,
            "rows": rows,
        }
    )


# --------------------------------------------------------------------------
# exports — the browser sends back the rows it is showing, edits included
# --------------------------------------------------------------------------

def _attachment(buf: io.BytesIO, media_type: str, filename: str) -> StreamingResponse:
    return StreamingResponse(
        buf,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/export/xlsx")
async def export_xlsx(payload: ExportRequest):
    rows = payload.rows
    if not rows:
        raise HTTPException(400, "Nothing to export.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Redirects"

    headers = ["404 URL", "200 redirect", "Confidence", "Match type", "Matched on", "Notes"]
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor="16202B")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    bands = {
        "strong": PatternFill("solid", fgColor="DFF3E8"),
        "fair": PatternFill("solid", fgColor="FBF0D6"),
        "weak": PatternFill("solid", fgColor="FADFDC"),
    }

    for r in rows:
        ws.append([r.dead, r.live, r.score, r.method, r.matched_on, r.notes])
        i = ws.max_row
        key = "strong" if r.score >= 85 else "fair" if r.score >= 60 else "weak"
        ws.cell(row=i, column=3).fill = bands[key]
        ws.cell(row=i, column=3).number_format = "0"
        for col in range(1, len(headers) + 1):
            ws.cell(row=i, column=col).font = Font(name="Arial")

    for col, width in zip(range(1, len(headers) + 1), (78, 78, 12, 18, 26, 30)):
        ws.column_dimensions[get_column_letter(col)].width = width

    legend = wb.create_sheet("How to read this")
    legend["A1"], legend["B1"] = "Confidence", "What it means"
    legend["A1"].font = Font(name="Arial", bold=True)
    legend["B1"].font = Font(name="Arial", bold=True)
    guide = [
        ("85-100", "Same page. The slug matched outright, or every distinctive word lines up. Safe to publish."),
        ("60-84", "Very likely the same page, worded slightly differently. Skim before publishing."),
        ("0-59", "Best available destination, not a confident match. Check these by hand."),
    ]
    for i, (band, meaning) in enumerate(guide, start=2):
        legend[f"A{i}"] = band
        legend[f"B{i}"] = meaning
        legend[f"A{i}"].font = Font(name="Arial")
        legend[f"B{i}"].font = Font(name="Arial")
    legend.column_dimensions["A"].width = 14
    legend.column_dimensions["B"].width = 96

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return _attachment(
        buf,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "redirects.xlsx",
    )


@app.post("/api/export/csv")
async def export_csv(payload: ExportRequest):
    if not payload.rows:
        raise HTTPException(400, "Nothing to export.")
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["404 URL", "200 redirect", "Confidence", "Match type", "Matched on", "Notes"])
    for r in payload.rows:
        writer.writerow([r.dead, r.live, round(r.score), r.method, r.matched_on, r.notes])
    return _attachment(io.BytesIO(out.getvalue().encode()), "text/csv", "redirects.csv")


@app.post("/api/export/rules")
async def export_rules(payload: ExportRequest, flavour: str = "nginx"):
    if not payload.rows:
        raise HTTPException(400, "Nothing to export.")
    out = io.StringIO()

    if flavour == "htaccess":
        out.write("# 301 redirects - generated by Redirect Mapper\n")
        for r in payload.rows:
            if not r.live:
                continue
            src = urlparse(r.dead)
            if src.query:
                continue  # query-string sources need RewriteCond, not a plain Redirect
            out.write(f"Redirect 301 {src.path or '/'} {r.live}\n")
        name = "htaccess.txt"
    else:
        out.write("# map $request_uri $redirect_target { ... }  - generated by Redirect Mapper\n")
        for r in payload.rows:
            if not r.live:
                continue
            src = urlparse(r.dead)
            path = src.path or "/"
            if src.query:
                path = f"{path}?{src.query}"
            out.write(f'"{path}" "{r.live}";\n')
        name = "nginx-map.conf"

    return _attachment(io.BytesIO(out.getvalue().encode()), "text/plain", name)


@app.get("/")
async def home():
    return FileResponse(STATIC_DIR / "index.html")
